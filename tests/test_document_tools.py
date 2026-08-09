import io

import fitz
from docx import Document
from openpyxl import Workbook

from backend.app.document_tools import create_pdf, docx_to_text, merge_pdfs, xlsx_to_text


def test_docx_extracts_text():
    doc = Document(); doc.add_paragraph("Akashic document test")
    buf = io.BytesIO(); doc.save(buf)
    assert "Akashic document test" in docx_to_text(buf.getvalue())


def test_xlsx_extracts_cells():
    wb = Workbook(); ws = wb.active; ws.append(["name", "roll"]); ws.append(["Jane", 123])
    buf = io.BytesIO(); wb.save(buf)
    text = xlsx_to_text(buf.getvalue())
    assert "Jane" in text and "123" in text


def test_pdf_create_and_merge():
    first = create_pdf("First", "one")
    second = create_pdf("Second", "two")
    from backend.app.artifacts import get_artifact
    _, a = get_artifact(first["artifact_id"])
    _, b = get_artifact(second["artifact_id"])
    merged = merge_pdfs([a, b])
    _, data = get_artifact(merged["artifact_id"])
    pdf = fitz.open(stream=data, filetype="pdf")
    assert pdf.page_count >= 2
